"""In-process text extraction for uploaded chat documents."""

from __future__ import annotations

import io
from pathlib import PurePath

from legalos_common.logging import get_logger

from app.infrastructure.image_vision import describe_image_with_gemini

logger = get_logger(__name__)

EXTRACT_CHAR_CAP = 200_000

_TEXT_EXT = {".txt", ".md", ".csv", ".json"}
_IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp"}
_IMAGE_CTYPES = {"image/jpeg", "image/png", "image/webp"}


def extract_document_text(
    data: bytes,
    *,
    filename: str | None = None,
    content_type: str | None = None,
) -> tuple[str, int]:
    """Return (text, page_count). page_count is 1 for non-paginated formats."""
    name = (filename or "").lower()
    ctype = (content_type or "").lower()
    suffix = PurePath(name).suffix

    if suffix in _TEXT_EXT or ctype.startswith("text/") or ctype in {"application/json", "text/csv"}:
        return _decode_text(data), 1
    if suffix == ".pdf" or ctype == "application/pdf":
        return _extract_pdf(data)
    if suffix == ".docx" or ctype in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }:
        return _extract_docx(data)
    if suffix == ".doc":
        text = _extract_docx(data)
        if text[0]:
            return text
        return _decode_text(data), 1
    if suffix in _IMAGE_EXT or ctype in _IMAGE_CTYPES or ctype.startswith("image/"):
        return _extract_image(data, filename=filename, content_type=content_type)
    return _decode_text(data), 1


def _decode_text(data: bytes) -> str:
    try:
        raw = data.decode("utf-8")
    except UnicodeDecodeError:
        raw = data.decode("latin-1", errors="ignore")
    return raw.strip()[:EXTRACT_CHAR_CAP]


def _extract_pdf(data: bytes) -> tuple[str, int]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        combined = "\n\n".join(pages).strip()[:EXTRACT_CHAR_CAP]
        return combined, len(reader.pages)
    except Exception as exc:
        logger.warning("pdf_extract_failed error=%s", exc)
        fallback = _decode_text(data)
        return fallback, 1


def _extract_docx(data: bytes) -> tuple[str, int]:
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()[:EXTRACT_CHAR_CAP], 1
    except Exception as exc:
        logger.warning("docx_extract_failed error=%s", exc)
        return "", 1


def _prepare_image(data: bytes):
    from PIL import Image, ImageOps

    img = Image.open(io.BytesIO(data))
    img = ImageOps.exif_transpose(img)
    if img.mode not in {"RGB", "L"}:
        img = img.convert("RGB")
    return img


def _extract_image(
    data: bytes,
    *,
    filename: str | None = None,
    content_type: str | None = None,
) -> tuple[str, int]:
    """OCR text from image bytes, with Gemini vision fallback when OCR is empty."""
    text = ""
    try:
        img = _prepare_image(data)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=90)
        prepared = buf.getvalue()

        try:
            import pytesseract

            text = pytesseract.image_to_string(img).strip()
        except Exception as exc:
            logger.warning("image_ocr_failed error=%s", exc)

        if not text:
            text = describe_image_with_gemini(
                prepared,
                filename=filename,
                content_type=content_type or "image/jpeg",
            )
    except Exception as exc:
        logger.warning("image_prepare_failed error=%s", exc)
        text = describe_image_with_gemini(data, filename=filename, content_type=content_type)

    return text[:EXTRACT_CHAR_CAP], 1
